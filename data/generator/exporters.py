"""Write the exports, each in its own (deliberately imperfect) format."""

from __future__ import annotations

import csv
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .config import Config
from .locations import Locations
from .model import Signal, Stay
from .rng import Rng

MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
TRAVELER_TYPES = {"Couple": 0.45, "Family with young children": 0.25, "Family with older children": 0.1, "Group of friends": 0.08, "Solo traveller": 0.07, "Business": 0.05}
TRIP_TYPES = {"Couples": 0.45, "Family": 0.35, "Friends": 0.12, "Solo": 0.05, "Business": 0.03}


def assign_ids(signals: list[Signal], rng: Rng) -> None:
    r = rng.spawn("ids")
    wos = sorted((s for s in signals if s.source == "work_order"), key=lambda s: (s.ts, s.uid))
    n = 4012
    for s in wos:
        s.record_id, s.export_id = str(n), f"WO-{n}"
        n += r.weighted({1: 0.9, 2: 0.07, 3: 0.02, 5: 0.01})
    fds = sorted((s for s in signals if s.source == "front_desk"), key=lambda s: (s.ts, s.uid))
    for i, s in enumerate(fds, start=1):
        s.record_id, s.export_id = str(i), f"FD-{i}"
    seen: set[str] = set()
    for s in sorted((s for s in signals if s.source == "survey"), key=lambda s: (s.ts, s.uid)):
        rid = "".join(r.choice("0123456789abcdef") for _ in range(8))
        while rid in seen:
            rid = "".join(r.choice("0123456789abcdef") for _ in range(8))
        seen.add(rid)
        s.record_id, s.export_id = rid, f"SV-{rid}"
    for s in sorted((s for s in signals if s.source == "review"), key=lambda s: (s.ts, s.uid)):
        p = s.extra["platform"]
        if p == "google":
            rid = "g-" + "".join(r.choice("0123456789abcdef") for _ in range(10))
        elif p == "booking":
            rid = "b-" + str(r.randint(100000000, 999999999))
        else:
            rid = "ta-" + str(r.randint(800000000, 999999999))
        while rid in seen:
            rid = rid[:-1] + r.choice("0123456789")
        seen.add(rid)
        s.record_id, s.export_id = rid, f"RV-{rid}"


# ---------------------------------------------------------------------------------------
def write_work_orders(path: Path, signals: list[Signal]) -> int:
    rows = sorted((s for s in signals if s.source == "work_order"), key=lambda s: (s.ts, s.uid))
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";", lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
        w.writerow(["Ticket No", "Created", "Closed", "Location", "Category", "Description", "Closure Note", "Technician", "Status", "Priority", "Reported By", "Reopened"])
        for s in rows:
            e = s.extra
            closed = e["closed_at"].strftime("%d/%m/%Y %H:%M") if e["closed_at"] else ""
            w.writerow([s.record_id, s.ts.strftime("%d/%m/%Y %H:%M"), closed, s.location_text, e["category"], s.text, e["closure"], e["technician"], e["status"], e["priority"], e["reported_by"], e["reopened"]])
    return len(rows)


def write_front_desk(path: Path, signals: list[Signal], cfg: Config) -> int:
    rows = sorted((s for s in signals if s.source == "front_desk"), key=lambda s: (s.ts, s.uid))
    wb = Workbook()
    wb.remove(wb.active)
    by_month: dict[tuple[int, int], list[Signal]] = {}
    for s in rows:
        by_month.setdefault((s.ts.year, s.ts.month), []).append(s)
    header = ["#", "Date", "Time", "Room", "Guest comment (original)", "Agent note", "Dept", "Resolution", "Shift"]
    for (y, mth), lst in sorted(by_month.items()):
        ws = wb.create_sheet(title=f"{MONTHS[mth - 1]}")
        ws["A1"] = f"Guest complaints & requests log - {cfg.short_name} - {MONTHS[mth - 1]} {y}"
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([])
        ws.append(header)
        for c in range(1, len(header) + 1):
            ws.cell(row=3, column=c).font = Font(bold=True)
        for s in lst:
            e = s.extra
            ws.append([int(s.record_id), s.ts.date(), s.ts.strftime("%H:%M"), e["room_cell"], s.text, e["agent_note"], e["department"], e["resolution"], e["shift"]])
        for col, width in zip("ABCDEFGHI", [6, 12, 7, 10, 70, 45, 14, 30, 7]):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=4, min_col=2, max_col=2):
            for cell in row:
                cell.number_format = "DD/MM/YYYY"
    wb.save(path)
    return len(rows)


def write_surveys(path: Path, signals: list[Signal], stays_by_id: dict[int, Stay], rng: Rng) -> int:
    r = rng.spawn("surveys-export")
    rows = sorted((s for s in signals if s.source == "survey"), key=lambda s: (s.ts, s.uid))
    cols = ["response_id", "reservation_no", "checkout_date", "language", "overall", "room", "cleanliness", "bathroom", "food_beverage", "staff", "pool_beach", "spa", "value_for_money", "nps", "comment", "would_recommend"]
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(cols)
        for s in rows:
            sc = s.extra["scores"]
            overall = sc["overall"]
            nps = max(0, min(10, overall + r.randint(-1, 1)))
            rec = "Yes" if nps >= 8 else ("No" if nps <= 5 else r.choice(["Yes", "No", ""]))
            w.writerow([s.record_id, s.stay_id, s.extra["checkout"].isoformat(), s.language if s.text else "", overall, sc["room"], sc["cleanliness"], sc["bathroom"], sc["food_beverage"], sc["staff"], sc["pool_beach"], "" if sc["spa"] is None else sc["spa"], sc["value_for_money"], nps, s.text, rec])
    return len(rows)


def write_reviews(path: Path, signals: list[Signal], stays_by_id: dict[int, Stay], loc: Locations, cfg: Config, rng: Rng) -> int:
    r = rng.spawn("reviews-export")
    rows = sorted((s for s in signals if s.source == "review"), key=lambda s: (s.ts, s.uid))
    out = []
    for s in rows:
        e = s.extra
        stay = stays_by_id.get(s.stay_id)
        room_type = loc.room_by_code[stay.room].room_type if stay else ""
        if e["platform"] == "google":
            out.append({"platform": "google", "review_id": s.record_id, "date": s.ts.strftime("%Y-%m-%dT%H:%M:%SZ"), "rating": e["rating5"], "language": s.language, "text": e["title"] + ". " + s.text if r.chance(0.3) else s.text, "owner_reply": None})
        elif e["platform"] == "booking":
            out.append({"platform": "booking", "review_id": s.record_id, "date": s.ts.date().isoformat(), "score": e["rating10"], "title": e["title"], "positive": e["positive"], "negative": e["negative"], "language": s.language, "traveler_type": r.weighted(TRAVELER_TYPES), "room_type_stayed": room_type if r.chance(0.8) else "", "nights": stay.nights if stay else None, "stay_month": e["checkout"].strftime("%B %Y")})
        else:
            out.append({"platform": "tripadvisor", "review_id": s.record_id, "published": s.ts.date().isoformat(), "date_of_stay": e["checkout"].strftime("%B %Y"), "rating": e["rating5"], "title": e["title"], "text": s.text, "language": s.language, "trip_type": r.weighted(TRIP_TYPES)})
    payload = {"property": cfg.resort_name, "source": "review aggregator export (all platforms)", "count": len(out), "reviews": out}
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
    return len(out)


def write_master_data(out: Path, loc: Locations, stays: list[Stay], occupancy_rows: list[dict]) -> dict[str, int]:
    counts = {}
    with open(out / "rooms.csv", "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["room_code", "building", "building_name", "floor", "room_type", "view", "is_villa"])
        for rm in loc.rooms:
            w.writerow([rm.code, rm.building, loc.building_names[rm.building], rm.floor, rm.room_type, rm.view, "Y" if rm.is_villa else "N"])
    counts["rooms.csv"] = len(loc.rooms)
    with open(out / "facilities.csv", "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["facility_code", "name", "type", "area", "notes"])
        for f in loc.facilities:
            w.writerow([f.code, f.name, f.type, f.area, f.notes])
    counts["facilities.csv"] = len(loc.facilities)
    with open(out / "equipment.csv", "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["equipment_id", "name", "type", "building", "location", "serves", "installed_year"])
        for e in loc.equipment:
            w.writerow([e.id, e.name, e.type, e.building or "", e.location_text, e.serves, e.installed])
    counts["equipment.csv"] = len(loc.equipment)
    with open(out / "stays.csv", "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["reservation_no", "room", "check_in", "check_out", "nights", "adults", "children", "nationality", "market_segment", "board", "booked_on"])
        for s in stays:
            w.writerow([s.stay_id, s.room, s.check_in.isoformat(), s.check_out.isoformat(), s.nights, s.adults, s.children, s.nationality, s.segment, s.board, s.booked_on.isoformat()])
    counts["stays.csv"] = len(stays)
    with open(out / "occupancy_daily.csv", "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(occupancy_rows[0].keys()), lineterminator="\n")
        w.writeheader()
        w.writerows(occupancy_rows)
    counts["occupancy_daily.csv"] = len(occupancy_rows)
    return counts
